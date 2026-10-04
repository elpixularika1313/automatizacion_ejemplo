import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

def optimizar_inventario(archivo_entrada, archivo_salida):
    # 1. Cargar el Excel en bruto
    wb = openpyxl.load_workbook(archivo_entrada)
    ws = wb.active
    
    # 2. Configurar nuevas cabeceras
    ws['F1'] = "Valor Total"
    ws['G1'] = "Estado"
    
    # Estilos para alertas de stock
    rojo_alerta = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    texto_rojo = Font(color=Color(rgb="009C0006"), bold=True)
    
    # 3. Procesar fila por fila
    for row in range(2, ws.max_row + 1):
        stock_actual = ws[f'C{row}'].value
        stock_minimo = ws[f'D{row}'].value
        precio = ws[f'E{row}'].value
        
        # Calcular el Valor Total (Stock * Precio)
        if stock_actual is not None and precio is not None:
            ws[f'F{row}'] = stock_actual * precio
            ws[f'F{row}'].number_format = '$#,##0.00'
            ws[f'E{row}'].number_format = '$#,##0.00'
            
        # Lógica de Alerta de Stock
        if stock_actual is not None and stock_minimo is not None:
            if stock_actual < stock_minimo:
                ws[f'G{row}'] = "REABASTECER"
                for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
                    ws[f'{col}{row}'].fill = rojo_alerta
                    ws[f'{col}{row}'].font = texto_rojo
            else:
                ws[f'G{row}'] = "OK"

    # 4. Formato para las Cabeceras
    fondo_cabecera = PatternFill(start_color="FF4F81BD", end_color="FF4F81BD", fill_type="solid")
    fuente_cabecera = Font(color=Color(rgb="00FFFFFF"), bold=True)
    
    columnas = ['A', 'B', 'C', 'D', 'E', 'F', 'G']
    for col in columnas:
        ws[f'{col}1'].fill = fondo_cabecera
        ws[f'{col}1'].font = fuente_cabecera
        ws.column_dimensions[col].width = 16

    # 5. Guardar el archivo final
    wb.save(archivo_salida)
    print(f"Reporte automatizado guardado con éxito como: {archivo_salida}")

if __name__ == "__main__":
    optimizar_inventario("inventario_bruto.xlsx", "reporte_bodega_listo.xlsx")